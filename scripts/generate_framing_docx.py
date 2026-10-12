import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def build_content_sections(doc, is_compact=True):
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(9.5 if is_compact else 10)
    normal_style.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    normal_style.paragraph_format.line_spacing = 1.12 if is_compact else 1.15
    normal_style.paragraph_format.space_after = Pt(2 if is_compact else 3)

    def add_section_header(num_str, title_text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5 if is_compact else 7)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r_num = p.add_run(f"{num_str}. ")
        r_num.bold = True
        r_num.font.size = Pt(10.5 if is_compact else 11)
        r_num.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        r_txt = p.add_run(title_text)
        r_txt.bold = True
        r_txt.font.size = Pt(10.5 if is_compact else 11)
        r_txt.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    # 1. Identitas
    add_section_header("1", "Identitas Proyek")
    tbl_id = doc.add_table(rows=4, cols=2)
    tbl_id.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_id.autofit = False
    
    col_widths = [Inches(1.8), Inches(4.9)]
    id_data = [
        ("Nama Proyek", "SkillGap.AI (Sistem Pemetaan Kesenjangan Kompetensi & Rekomendasi Karier/Sertifikasi Berbasis AI)"),
        ("Sektor", "Education Technology (EdTech) & Talent Development / Career Analytics"),
        ("Nama Anggota Tim", "1. Julia Rahma Nurmawati (V3925009)\n2. Aisyah Nurul Ilmi Prianto (V3925018)"),
        ("Repository & URL Live", "Github: https://github.com/aisyahnarl/skillgap-ai\nURL Live: https://skillgap-ai-seven.vercel.app/ (Alternatif: https://skillgap-ai-kappa.vercel.app/)")
    ]
    for row_idx, (label, val) in enumerate(id_data):
        row = tbl_id.rows[row_idx]
        prevent_row_split(row)
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = col_widths[0], col_widths[1]
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF" if row_idx % 2 == 0 else "F8FAFC")
        set_cell_margins(c0, top=40, bottom=40, left=80, right=80)
        set_cell_margins(c1, top=40, bottom=40, left=80, right=80)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(1)
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.size = Pt(8.5 if is_compact else 9)
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(1)
        r1 = p1.add_run(val)
        r1.font.size = Pt(8.5 if is_compact else 9)

    # 2. Masalah
    add_section_header("2", "Masalah (Problem Framing)")
    p_m1 = doc.add_paragraph()
    r = p_m1.add_run("• Pengguna Sasaran: ")
    r.bold = True
    p_m1.add_run("Pelajar SMK Vokasi (RPL/TKJ), mahasiswa D3/S1 Vokasi IT, dan fresh graduate / career switcher bidang IT.")
    
    p_m2 = doc.add_paragraph()
    r = p_m2.add_run("• Masalah Utama: ")
    r.bold = True
    p_m2.add_run("Adanya kesenjangan kompetensi (skill gap) signifikan antara kurikulum formal dengan kebutuhan teknis industri IT. Calon talenta mengalami disorientasi kompetensi (tidak tahu apa yang kurang), kebingungan memilih sertifikasi yang kredibel dan efisien di antara ribuan sertifikasi beredar, serta ketiadaan panduan roadmap belajar bertahap yang terarah.")
    
    p_m3 = doc.add_paragraph()
    r = p_m3.add_run("• Mengapa Penting: ")
    r.bold = True
    p_m3.add_run("Tingginya angka pengangguran terbuka lulusan vokasi akibat mismatch keterampilan (BPS & WEF Report). Solusi diagnosis AI mencegah pemborosan biaya jutaan rupiah dan waktu berbulan-bulan untuk kursus/sertifikasi yang tidak relevan.")

    p_m4 = doc.add_paragraph()
    r = p_m4.add_run("• Karakteristik Masalah: ")
    r.bold = True
    p_m4.add_run("Spesifik (domain 20 peran kerja IT standar industri), Terukur (skor kecocokan Match Score 0-100% dan Gap Score kuantitatif per kompetensi), dan Terbatas (jenjang Pemula-Lanjutan dengan kurasi sertifikasi resmi BNSP, Dicoding, Google, AWS, Microsoft, Meta).")

    # 3. Input dan Output
    add_section_header("3", "Input dan Output Sistem")
    p_io1 = doc.add_paragraph()
    r = p_io1.add_run("• Data Input: ")
    r.bold = True
    p_io1.add_run("Profil pendidikan (Pelajar/Mahasiswa), target peran kerja (misal: Frontend Dev, Data Analyst), penilaian mandiri keterampilan (skala 5 level kemahiran: Beginner-Expert), respons kuesioner kesiapan belajar, dan teks deskripsi pengalaman/portofolio bebas (unstructured text).")

    p_io2 = doc.add_paragraph()
    r = p_io2.add_run("• Hasil Output: ")
    r.bold = True
    p_io2.add_run("Skor Kecocokan (Match Score %) dan Skor Kesenjangan (Skill Gap %), Matriks Status Keahlian (Strong, Developing, Gap), Rekomendasi Sertifikasi Berbobot (prioritas High/Medium, estimasi waktu, prasyarat), dan Roadmap Pembelajaran Personal 3 Fase (Fondasi, Latihan Terarah & Proyek, Validasi & Sertifikasi).")

    # 4. Tipe Masalah AI
    add_section_header("4", "Tipe Masalah AI")
    p_tp = doc.add_paragraph()
    r = p_tp.add_run("Tipe Masalah: ")
    r.bold = True
    p_tp.add_run("Generative AI (Structured Information Extraction & Retrieval-Augmented Generation / RAG) dikombinasikan dengan Machine Learning (Multi-Label Skill Matching & Vector Similarity Ranking).")
    
    p_tp_det = doc.add_paragraph()
    p_tp_det.add_run("1) Structured Extraction (Gen AI/NLP): Mengekstraksi keahlian dari teks portofolio pengguna ke format JSON terstandarisasi. 2) RAG Architecture: Menyematkan standar SKKNI dan silabus sertifikasi resmi sebagai konteks acuan untuk mencegah halusinasi. 3) Vector Cosine Similarity (ML): Menghitung kedekatan vektor antara kemampuan pengguna dan spesifikasi role target. 4) Contextual Narrative Generation (Gen AI): Menghasilkan penjelasan rekomendasi dan roadmap adaptif.")

    # 5. Justifikasi AI
    add_section_header("5", "Justifikasi AI (Mengapa Butuh AI)")
    p_just = doc.add_paragraph()
    p_just.add_run("• ")
    r = p_just.add_run("Variasi Linguistik Terminologi: ")
    r.bold = True
    p_just.add_run("Pengguna menyebutkan keahlian dengan istilah bahasa alami yang sangat beragam ('slicing Figma ke CSS', 'bikin API Node', 'ReactJS'). Logika if-else statis gagal menangani variasi sinonim semantik, sedangkan AI/NLP mampu memetakan variasi teks ke taksonomi baku.\n• ")
    r = p_just.add_run("Penalaran Kurikulum yang Adaptif: ")
    r.bold = True
    p_just.add_run("Rule-based hanya mampu memberikan template kaku, sedangkan Gen AI mampu menyesuaikan penalaran kurikulum berdasarkan persona (siswa SMK fokus fondasi konsep, mahasiswa fokus arsitektur perangkat lunak dan portofolio).\n• ")
    r = p_just.add_run("Skalabilitas & Kemudahan Pembaruan: ")
    r.bold = True
    p_just.add_run("Standar industri dan sertifikasi baru dapat ditambahkan ke basis data pengetahuan RAG tanpa merombak ribuan baris kode logika program.")

    # 6. Algoritma yang Dipakai
    add_section_header("6", "Algoritma yang Dipakai dan Alasannya")
    p_alg1 = doc.add_paragraph()
    r = p_alg1.add_run("1. Retrieval-Augmented Generation (RAG) + LLM (Gemini / Llama): ")
    r.bold = True
    p_alg1.add_run("Mencegah halusinasi fakta sertifikasi/SKKNI karena respon model terikat (grounded) pada basis data resmi, sekaligus memproduksi penalaran roadmap yang terstruktur rapi (JSON schema).\n")
    r = p_alg1.add_run("2. Vector Space Embedding & Cosine Similarity: ")
    r.bold = True
    p_alg1.add_run("Menghitung skor kecocokan (Match Score) secara matematis, cepat, deterministik, dan dapat diaudit (interpretable) berdasarkan vektor representasi kompetensi pengguna vs persyaratan peran kerja.")

    # 7. Dataset yang Digunakan
    add_section_header("7", "Dataset yang Digunakan")
    tbl_ds = doc.add_table(rows=6, cols=3)
    tbl_ds.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_ds.autofit = False
    ds_widths = [Inches(1.8), Inches(1.5), Inches(3.4)]
    ds_headers = ["Nama Dataset", "Sumber Data", "Tautan Resmi & Keterangan"]
    
    hdr_row = tbl_ds.rows[0]
    prevent_row_split(hdr_row)
    for c_idx, h_text in enumerate(ds_headers):
        cell = hdr_row.cells[c_idx]
        cell.width = ds_widths[c_idx]
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, top=50, bottom=50, left=70, right=70)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(h_text)
        r.bold = True
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    ds_rows_data = [
        ("IT Job Roles Skills Dataset (Primer)", "Kaggle\n(Dhivya Dharuna)", "https://www.kaggle.com/datasets/dhivyadharuna/it-job-roles-skills-dataset\nKatalog 493 role IT, deskripsi pekerjaan, skills, dan sertifikasi."),
        ("Candidate Job Role Dataset (Pendukung)", "Kaggle\n(C. K. Shetty)", "https://www.kaggle.com/datasets/ckshetty/candidate-job-role-dataset\nData profil kandidat dengan pengalaman, skills, dan target role."),
        ("Student Skill Gap Analysis (Pelajar/Mhs)", "Kaggle\n(Shristi Saini)", "https://www.kaggle.com/datasets/shristisaini/student-skill-gap-analysis\nDataset evaluasi disparitas skill mahasiswa/pelajar vs industri."),
        ("Resume Skill Gap Analyzer (NLP/Matching)", "Kaggle\n(Keerthi Ramesh)", "https://www.kaggle.com/datasets/keerthiramesh470/resume-skill-gap-analyzer-dataset\nData ekstraksi dan analisis kesenjangan skill berbasis resume."),
        ("SkillGap.AI Golden Test Set (Wajib Note b)", "Data Uji Internal\n(Terkurasi Manual)", "datasets/skillgap_ground_truth_test_set.csv\n30 kasus uji beranotasi kunci jawaban manual untuk evaluasi LLM.")
    ]

    for r_idx, r_data in enumerate(ds_rows_data, start=1):
        row = tbl_ds.rows[r_idx]
        prevent_row_split(row)
        bg = "FFFFFF" if r_idx % 2 != 0 else "F8FAFC"
        for c_idx, val in enumerate(r_data):
            cell = row.cells[c_idx]
            cell.width = ds_widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=35, bottom=35, left=60, right=60)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(val)
            r.font.size = Pt(7.5 if is_compact else 8)

    # 8. Ruang Lingkup
    add_section_header("8", "Ruang Lingkup Proyek (Scope of Work)")
    p_sc1 = doc.add_paragraph()
    r = p_sc1.add_run("• Yang Dikerjakan: ")
    r.bold = True
    p_sc1.add_run("1) Modul ekstraksi terstruktur keahlian dari teks input pengguna. 2) Algoritma penghitungan selisih gap kompetensi & Match Score berbasis kesamaan vektor. 3) Modul rekomendasi sertifikasi berbobot & generator roadmap belajar personal 3 fase berbasis RAG. 4) Pembuatan dataset set uji (30 kasus uji) dan evaluasi performa model kuantitatif. 5) Integrasi antarmuka web interaktif (Next.js / Vite) untuk asesmen pengguna.")

    p_sc2 = doc.add_paragraph()
    r = p_sc2.add_run("• Yang TIDAK Dikerjakan: ")
    r.bold = True
    p_sc2.add_run("1) Scraping portal lowongan eksternal secara real-time saat runtime aplikasi berjalan. 2) Penyelenggaraan ujian sertifikasi dan penerbitan sertifikat resmi (sistem hanya bertindak sebagai pemeta persiapan). 3) Melakukan pre-training model bahasa besar (LLM ratusan miliar parameter) dari awal, melainkan memanfaatkan foundation model via API dengan integrasi RAG.")

    # Box Note Evaluasi Gen AI
    tbl_note = doc.add_table(rows=1, cols=1)
    tbl_note.alignment = WD_TABLE_ALIGNMENT.CENTER
    prevent_row_split(tbl_note.rows[0])
    c_note = tbl_note.rows[0].cells[0]
    c_note.width = Inches(6.7)
    set_cell_background(c_note, "EFF6FF")
    set_cell_margins(c_note, top=50, bottom=50, left=100, right=100)
    p_note_content = c_note.paragraphs[0]
    p_note_content.paragraph_format.space_after = Pt(1)
    r_nh = p_note_content.add_run("Kepatuhan Catatan Dosen (Note b — Evaluasi Objektif Gen AI):\n")
    r_nh.bold = True
    r_nh.font.size = Pt(8.5)
    r_nh.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)
    r_nb = p_note_content.add_run(
        "Telah disiapkan set uji teranotasi manual (datasets/skillgap_ground_truth_test_set.csv) dengan 30 skenario uji. "
        "Evaluasi objektif model diukur menggunakan metrik: (1) F1-Score Ekstraksi Keahlian (Target >= 85%), "
        "(2) Jaccard Similarity Deteksi Gap (Target >= 80%), (3) Mean Absolute Error (MAE) Match Score (Target <= 7.5%), "
        "dan (4) Top-1 Accuracy Rekomendasi Sertifikasi (Target >= 85%)."
    )
    r_nb.font.size = Pt(8)
    r_nb.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

def create_compact_version(output_path):
    doc = docx.Document()
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("DOKUMEN PROJECT FRAMING")
    r_title.bold = True
    r_title.font.size = Pt(14)
    r_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Mata Kuliah Blok Kecerdasan Buatan & Rekayasa Perangkat Lunak\nProgram Studi D3 Teknik Informatika Kampus Madiun — Sekolah Vokasi UNS")
    r_sub.font.size = Pt(9)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    p_sub.paragraph_format.space_after = Pt(6)

    build_content_sections(doc, is_compact=True)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Compact version saved to: {output_path}")

def create_cover_version(output_path):
    doc = docx.Document()
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    # COVER PAGE
    p_c1 = doc.add_paragraph()
    p_c1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c1.paragraph_format.space_before = Pt(36)
    p_c1.paragraph_format.space_after = Pt(12)
    r_c1 = p_c1.add_run("LAPORAN PROJECT FRAMING")
    r_c1.bold = True
    r_c1.font.size = Pt(16)
    r_c1.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p_c2 = doc.add_paragraph()
    p_c2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c2.paragraph_format.space_after = Pt(120)
    r_c2 = p_c2.add_run("SkillGap.AI : Sistem Pemetaan Kesenjangan Kompetensi dan Rekomendasi\nPengembangan Diri, Persiapan Studi Lanjut, Karier, serta Sertifikasi\nBerbasis AI")
    r_c2.bold = True
    r_c2.font.size = Pt(12)
    r_c2.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    p_c3 = doc.add_paragraph()
    p_c3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c3.paragraph_format.space_after = Pt(10)
    r_c3 = p_c3.add_run("Disusun Oleh :")
    r_c3.font.size = Pt(10.5)

    p_c4 = doc.add_paragraph()
    p_c4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c4.paragraph_format.space_after = Pt(140)
    r_c4_1 = p_c4.add_run("JULIA RAHMA NURMAWATI (V3925009)\n")
    r_c4_1.bold = True
    r_c4_1.font.size = Pt(11)
    r_c4_2 = p_c4.add_run("AISYAH NURUL ILMI PRIANTO (V3925018)")
    r_c4_2.bold = True
    r_c4_2.font.size = Pt(11)

    p_c5 = doc.add_paragraph()
    p_c5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c5 = p_c5.add_run(
        "PROGRAM STUDI D3 TEKNIK INFORMATIKA\n"
        "KAMPUS KAB. MADIUN\n"
        "SEKOLAH VOKASI\n"
        "UNIVERSITAS SEBELAS MARET\n"
        "2026"
    )
    r_c5.bold = True
    r_c5.font.size = Pt(11)

    # Page break after cover
    doc.add_page_break()

    build_content_sections(doc, is_compact=False)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Cover version saved to: {output_path}")

if __name__ == "__main__":
    create_compact_version(os.path.join("docs", "DOKUMEN_PROJECT_FRAMING_SKILLGAP_AI_REVISI.docx"))
    create_cover_version(os.path.join("docs", "DOKUMEN_PROJECT_FRAMING_SKILLGAP_AI_DENGAN_COVER.docx"))
