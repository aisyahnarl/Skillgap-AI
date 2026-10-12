import json
import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=75, bottom=75, left=115, right=115):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_box_section(doc, title_str, body_text=None, bg_hex="F8FAFC", border_title_rgb=RGBColor(0x0F, 0x17, 0x2A), is_code=False, img_path=None):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)

    p0 = cell.paragraphs[0]
    p0.paragraph_format.space_after = Pt(3)
    r_title = p0.add_run(title_str)
    r_title.bold = True
    r_title.font.size = Pt(9.5)
    r_title.font.color.rgb = border_title_rgb

    if body_text:
        p1 = cell.add_paragraph()
        p1.paragraph_format.space_after = Pt(2)
        p1.paragraph_format.line_spacing = 1.06 if is_code else 1.14
        r_body = p1.add_run(body_text)
        r_body.font.name = 'Consolas' if is_code else 'Calibri'
        r_body.font.size = Pt(7.5 if is_code else 9.5)
        r_body.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    if img_path and os.path.exists(img_path):
        p_img = cell.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(5)
        p_img.paragraph_format.space_after = Pt(2)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(6.05))

    doc.add_paragraph().paragraph_format.space_after = Pt(3)

def create_word_report(res, output_docx="docs/LAPORAN_PRAKTIKUM_HARI1_SKILLGAP_AI.docx"):
    doc = docx.Document()

    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    normal_style.paragraph_format.line_spacing = 1.14
    normal_style.paragraph_format.space_after = Pt(4)

    # Top Banner
    tbl_banner = doc.add_table(rows=1, cols=1)
    tbl_banner.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_ban = tbl_banner.rows[0].cells[0]
    c_ban.width = Inches(6.5)
    set_cell_background(c_ban, "1D5B79")
    set_cell_margins(c_ban, top=65, bottom=65, left=110, right=110)
    p_ban = c_ban.paragraphs[0]
    r_ban = p_ban.add_run("PRAKTIKUM PYTHON UNTUK SISTEM CERDAS")
    r_ban.bold = True
    r_ban.font.size = Pt(10.5)
    r_ban.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    p_t = doc.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t.paragraph_format.space_before = Pt(8)
    p_t.paragraph_format.space_after = Pt(6)
    r_t = p_t.add_run(
        "MODUL PRAKTIKUM\n"
        "MATA KULIAH PYTHON UNTUK SISTEM CERDAS\n"
        "PROGRAM STUDI D3 TEKNIK INFORMATIKA K KAB. MADIUN\n"
        "SEKOLAH VOKASI\n"
        "UNIVERSITAS SEBELAS MARET"
    )
    r_t.bold = True
    r_t.font.size = Pt(11)
    r_t.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    # Tabel Identitas
    tbl_id = doc.add_table(rows=6, cols=2)
    tbl_id.alignment = WD_TABLE_ALIGNMENT.CENTER
    id_rows = [
        ("HARI 1", "FONDASI DATA DAN MACHINE LEARNING UNTUK PROJECT CAPSTONE"),
        ("NAMA MAHASISWA", "1. Julia Rahma Nurmawati\n2. Aisyah Nurul Ilmi Prianto"),
        ("NIM", "1. V3925009\n2. V3925018"),
        ("KELAS", "TI-B / Reguler Kampus Kab. Madiun"),
        ("ANGKATAN", "2025"),
        ("NAMA PROJECT / TIM", "SkillGap AI (No. 15 Tabel Panduan Tugas Mandiri per Project)\nDataset Project: DMA DATASET.xlsx (Student Skill Gap Analysis — 500 Baris, 15 Kolom)")
    ]
    for idx, (k, v) in enumerate(id_rows):
        row = tbl_id.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = Inches(1.9), Inches(4.6)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=55, bottom=55, left=90, right=90)
        set_cell_margins(c1, top=55, bottom=55, left=90, right=90)
        r0 = c0.paragraphs[0].add_run(k)
        r0.bold = True
        r0.font.size = Pt(9.5)
        r1 = c1.paragraphs[0].add_run(v)
        if idx == 0:
            r1.bold = True
        r1.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def add_main_heading(judul):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(judul)
        r.bold = True
        r.font.size = Pt(12.5)
        r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    def add_sub_heading(judul):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(judul)
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x1D, 0x5B, 0x79)

    # ==========================================================================
    # BAGIAN III: LANGKAH PRAKTIKUM (OUTPUT A.1 s.d. A.13 PADA DMA DATASET.xlsx)
    # ==========================================================================
    add_main_heading("III. HASIL LANGKAH PRAKTIKUM (A.1 – A.13 PADA 'DMA DATASET.xlsx')")

    add_sub_heading("A.1 Import Library dan Membaca Data ('DMA DATASET.xlsx')")
    add_box_section(doc, "Kode Python A.1:", "import time, numpy as np, pandas as pd, matplotlib.pyplot as plt, seaborn as sns\ndf = pd.read_excel('DMA DATASET.xlsx')\ndf.head()", bg_hex="F1F5F9", is_code=True)
    add_box_section(doc, "Output A.1 – Lima Baris Pertama Dataset (df.head()):", res["out_a1"], bg_hex="F8FAFC", is_code=True)

    add_sub_heading("A.2 Eksplorasi Awal dan Pemeriksaan Kualitas Data ('DMA DATASET.xlsx')")
    add_box_section(doc, "Output A.2.1 – Ukuran Data (500, 15) dan Informasi Tipe Kolom (df.info()):", res["out_a2_1"], bg_hex="F8FAFC", is_code=True)
    add_box_section(doc, "Output A.2.2 – Missing Value (Unnamed: 13 = 500), Duplikat, 9 Variasi Kolom 'Year ', dan Distribusi Label:", res["out_a2_2"], bg_hex="F8FAFC", is_code=True)
    add_box_section(doc, "Output A.2.3 – Statistik Deskriptif Numerik & Kategorikal (df.describe()):", res["out_a2_3"], bg_hex="F8FAFC", is_code=True)

    add_sub_heading("A.3 NumPy: Looping vs Vektorisasi")
    add_box_section(doc, "Output A.3 – Perbandingan Waktu Looping Python dan Vektorisasi NumPy (DMA DATASET):", res["out_a3"], bg_hex="F8FAFC", is_code=True)

    add_sub_heading("A.4 Visualisasi EDA ('DMA DATASET.xlsx')")
    add_box_section(doc, "Output A.4 – Empat Grafik EDA pada DMA DATASET.xlsx:", "1. Distribusi Minat Karier per Kategori Gap | 2. Inkonsistensi Penulisan Kolom 'Year ' (9 Variasi) | 3. Sebaran Skor Kesiapan vs Pengalaman Proyek | 4. Heatmap Korelasi Numerik", bg_hex="FFFFFF", img_path="docs/images/soal2_eda_skillgap.png")

    add_sub_heading("A.5 Pembersihan Data ('DMA DATASET.xlsx')")
    add_box_section(doc, "Output A.5 – Ukuran Data (500, 12), Penyeragaman 'Year', dan Penghapusan 'Unnamed: 13' & PII:", res["out_a5"], bg_hex="F8FAFC", is_code=True)

    add_sub_heading("A.6 Feature Engineering Berbasis Pengetahuan Domain")
    add_box_section(doc, "Output A.6 – Fitur Turunan (total_item_skill, skor_kesiapan_industri, indeks_defisit_skill) dan Label Numerik:", res["out_a6"], bg_hex="F8FAFC", is_code=True)

    add_sub_heading("A.7 Split Data dan Menyusun Pipeline Preprocessing")
    add_box_section(doc, "Output A.7 – Ukuran Data Latih (400, 11) dan Data Uji (100, 11):", res["out_a7"], bg_hex="F8FAFC", is_code=True)

    add_sub_heading("A.8 Membangun Baseline & A.9 Membandingkan Model dengan 5-Fold Cross-Validation")
    add_box_section(doc, "Output A.8 – Perbandingan Baseline Dummy, Rule-Based (Kata Kunci Manual), dan Pipeline ML:", res["out_a8"], bg_hex="F8FAFC", is_code=True)
    add_box_section(doc, "Output A.9 – F1-score Rata-rata dan Simpangan Tiap Model (5-Fold Cross-Validation):", res["out_a9"], bg_hex="F8FAFC", is_code=True)

    add_sub_heading("A.10 Simulasi Data Leakage")
    add_box_section(doc, "Output A.10 – Skor Model dengan Data Leakage dan Pemeriksaan Overlap Kandidat:", res["out_a10"], bg_hex="F8FAFC", is_code=True)

    add_sub_heading("A.11 Unsupervised Learning: TF-IDF + K-Means pada DMA DATASET & Isolation Forest")
    add_box_section(doc, "Output A.11.1 & A.11.2 – Silhouette Score (Optimal k=6), Kata Kunci Klaster Karier, dan Deteksi Anomali:", res["out_a11"], bg_hex="F8FAFC", is_code=True, img_path="docs/images/soal4_kmeans_baseline.png")

    add_sub_heading("A.12 Feature Engineering Teks (TF-IDF), Time Series (Lag/Rolling), dan Retrieval RAG")
    add_box_section(doc, "Output A.12.1 – Kata dengan Bobot TF-IDF Tertinggi per Kelas pada DMA DATASET:", res["out_a12_1"], bg_hex="F8FAFC", is_code=True)
    add_box_section(doc, "Output A.12.2 – Fitur Lag dan Rolling Mean Harian (shift(1)):", res["out_a12_2"], bg_hex="F8FAFC", is_code=True)
    add_box_section(doc, "Output A.12.3 – Hasil Retrieval TF-IDF Knowledge Base, Recall@1, Recall@3, dan MRR:", res["out_a12_3"], bg_hex="F8FAFC", is_code=True)

    add_sub_heading("A.13 Menyimpan Model dan Data Bersih")
    add_box_section(doc, "Output A.13 – Skor Model Final pada Data Uji dan Konfirmasi Penyimpanan:", res["out_a13"], bg_hex="F8FAFC", is_code=True)

    # ==========================================================================
    # BAGIAN IV: TUGAS MANDIRI (SOAL 1 s.d. SOAL 5 BERBASIS DMA DATASET.xlsx)
    # ==========================================================================
    add_main_heading("IV. TUGAS MANDIRI (SOAL 1 – SOAL 5 BERBASIS 'DMA DATASET.xlsx')")

    # SOAL 1
    add_sub_heading("Soal 1 – Project Framing")
    jawaban_soal_1 = (
        "a. Nama Project, Sektor, dan Pengguna Utama Sistem:\n"
        "   • Nama Project : SkillGap.AI (Sistem Pemetaan Kesenjangan Kompetensi & Rekomendasi Pengembangan Diri, Karier, serta Sertifikasi Berbasis AI).\n"
        "   • Sektor       : Education Technology (EdTech) & Human Capital / Career Analytics.\n"
        "   • Pengguna Utama: Mahasiswa & Pelajar Vokasi/Perguruan Tinggi (B.Tech IT, B.Tech CSE, B.Sc Data Science, M.Tech AI, BBA, MBA) serta pencari kerja pemula.\n\n"
        "b. Masalah yang Diselesaikan secara Spesifik, Terukur, dan Terbatas:\n"
        "   • Spesifik : Mengatasi disparitas kompetensi (skill gap) antara keterampilan teknis/softskill mahasiswa saat ini dengan target minat karier ('Career Interest') serta menentukan dukungan intervensi yang tepat ('Support required').\n"
        "   • Terukur  : Mengukur Rating Teknis (1–5), Rating Soft Skill (1–5), Skor Kesiapan Industri (0–100), mengelompokkan klaster kompetensi dengan Silhouette Score, dan mengklasifikasikan mahasiswa ke dalam kategori 'High_Gap' (butuh pelatihan intensif = 1) atau 'Ready' (siap kerja/magang = 0).\n"
        "   • Terbatas : Menggunakan 'DMA DATASET.xlsx' (500 mahasiswa, 15 kolom) yang mencakup 6 target karier utama ('Cloud Architect', 'Software Developer', 'Data Scientist', 'Cybersecurity Expert', 'Business Analyst', 'AI Engineer').\n\n"
        "c. Input dan Output Sistem:\n"
        "   • Input Sistem : Tahun kuliah ('Year'), program studi ('Current Course'), daftar keahlian teknis ('Technical Skills'), bahasa pemrograman ('Programming Languages'), nilai 'Rating' teknis (1–5), 'Soft Skills', nilai 'Rating.1' softskill (1–5), pengalaman proyek ('Projects': Yes/No), minat karier ('Career Interest'), tantangan ('Challenges'), dan metode belajar ('Method').\n"
        "   • Output Sistem: (1) Klasterisasi pemetaan keahlian & minat karier berbasis TF-IDF + K-Means, (2) Prediksi status kesenjangan ('High_Gap' vs 'Ready'), serta (3) Rekomendasi program intervensi & dokumen sertifikasi berbasis Retrieval RAG.\n\n"
        "d. Tipe Masalah AI:\n"
        "   • Sesuai Tabel Panduan No. 15 (SkillGap AI): Kombinasi NLP & Clustering (TF-IDF + K-Means pada teks skill & minat karier), Supervised Classification (deteksi 'High_Gap' vs 'Ready'), serta Generative AI / RAG.\n\n"
        "e. Alasan Masalah Ini Membutuhkan AI dan Baseline Non-AI Pembanding:\n"
        "   • Alasan Butuh AI: Kombinasi multivariat antara jurusan kuliah ('Current Course'), teks multi-label ('Technical Skills', 'Programming Languages', 'Soft Skills'), pengalaman proyek, dan rating tidak dapat dipisahkan dengan aturan kata kunci statis karena tingginya variasi kombinasi skill.\n"
        "   • Baseline Non-AI Pembanding (Sesuai Tabel Panduan No. 15): Pencocokan Daftar Kata Kunci Skill Manual (exact keyword matching antara Technical Skills dan kamus manual Career Interest).\n\n"
        "f. Metrik Evaluasi Utama beserta Alasannya:\n"
        "   • Classification: F1-Score dan Recall menjadi metrik utama karena distribusi label tidak seimbang (61.8% Ready vs 38.2% High_Gap) dan Recall tinggi krusial agar mahasiswa yang mengalami High_Gap tidak terlewat.\n"
        "   • Clustering: Silhouette Score (-1 s.d. +1) untuk memilih jumlah klaster k terbaik pada TF-IDF + K-Means.\n"
        "   • Retrieval RAG: Recall@1, Recall@3, dan Mean Reciprocal Rank (MRR).\n\n"
        "g. Risiko Bias dan Privasi Data beserta Mitigasinya:\n"
        "   • Risiko Privasi Data: 'DMA DATASET.xlsx' memuat identitas langsung berupa nama lengkap ('Name') dan alamat email ('email_id') dari 500 mahasiswa. Mitigasi: Menghapus kolom 'Name' dan 'email_id' (data anonymization) sebelum pemodelan.\n"
        "   • Risiko Bias: Potensi bias antar program studi ('MBA'/'BBA' vs 'B.Tech'). Mitigasi: Stratified splitting (stratify=y) dan evaluasi cross-validation 5-fold.\n\n"
        "h. Ruang Lingkup (Apa yang Dikerjakan dan TIDAK Dikerjakan dalam Blok 5 Hari):\n"
        "   • Yang Dikerjakan: EDA menyeluruh pada 'DMA DATASET.xlsx', pembersihan spasi nama kolom & 9 variasi penulisan 'Year ', penghapusan kolom kosong 'Unnamed: 13' & PII, pembuatan fitur turunan domain, TF-IDF + K-Means clustering, Pipeline Classification 5-Fold CV, Retrieval TF-IDF, dan pemeriksaan data leakage.\n"
        "   • Yang TIDAK Dikerjakan: Live scraping portal lowongan kerja eksternal saat runtime dan pre-training LLM dari nol."
    )
    add_box_section(doc, "Jawaban Soal 1 (Project Framing):", jawaban_soal_1, bg_hex="EFF6FF", border_title_rgb=RGBColor(0x1E, 0x3A, 0x8A))

    # SOAL 2 (FOKUS UTAMA PERMINTAAN USER: DMA DATASET.xlsx)
    add_sub_heading("Soal 2 – EDA Dataset Project ('DMA DATASET.xlsx')")
    p_s2_intro = doc.add_paragraph(
        "Eksplorasi Data (EDA) dilakukan secara langsung pada dataset project tim yaitu 'DMA DATASET.xlsx' "
        "(Student Skill Gap Analysis Dataset) yang berada di folder utama WEBSITE SKILLGAP."
    )
    add_box_section(
        doc,
        "Output Soal 2 (a, b, c) – Ukuran Data, Tipe Kolom (df.info()), Missing Value, Duplikat, Variasi Kategori, & Distribusi Label pada 'DMA DATASET.xlsx':",
        res["out_a2_1"] + "\n" + res["out_a2_2"] + "\n\n" + res["out_a2_3"],
        bg_hex="F8FAFC",
        is_code=True,
        img_path="docs/images/soal2_eda_skillgap.png"
    )

    insight_eda_text = (
        "Berdasarkan hasil eksplorasi (EDA) dan keempat grafik visualisasi pada 'DMA DATASET.xlsx' (500 baris, 15 kolom), diperoleh 4 insight utama beserta dampaknya bagi pemodelan:\n\n"
        "1. Insight 1 (Temuan Kolom 100% Missing Value 'Unnamed: 13' & Spasi Tersembunyi pada Nama Kolom):\n"
        "   • Temuan Data: Dari 15 kolom pada 'DMA DATASET.xlsx', kolom 'Unnamed: 13' memiliki 500 missing value (100% bernilai NaN / kosong), sedangkan 14 kolom lainnya memiliki 0 missing value dan 0 baris duplikat penuh. Selain itu, terdapat 5 nama kolom yang mengandung spasi tidak konsisten di awal/akhir ('Year ', '     Programming Languages', ' Projects', 'Challenges ', ' Method') serta nama kolom ambigu ('Rating' untuk Technical Skills dan 'Rating.1' untuk Soft Skills).\n"
        "   • Dampak bagi Pemodelan: Kolom 'Unnamed: 13' wajib dihapus (drop) karena tidak mengandung informasi apa pun dan akan menyebabkan kegagalan pada SimpleImputer. Seluruh nama kolom wajib dibersihkan dengan .str.strip() dan di-rename ('Technical_Rating', 'SoftSkill_Rating') agar pemanggilan fitur pada ColumnTransformer tidak memicu KeyError.\n\n"
        "2. Insight 2 (Inkonsistensi Format Penulisan pada Kolom Kategorik 'Year ' — 9 Variasi Penulisan):\n"
        "   • Temuan Data: Grafik 2 dan pemeriksaan df['Year '].value_counts() membuktikan adanya 9 variasi penulisan berbeda untuk 4 tingkat tahun perkuliahan, yaitu: 'Final Year' (129), 'Final' (2), 'Final ' (1), '2nd Year' (126), '2nd' (4), '2nd ' (2), '3rd Year' (124), '1st Year' (110), dan '1st ' (2).\n"
        "   • Dampak bagi Pemodelan: Apabila tidak diseragamkan terlebih dahulu, OneHotEncoder akan menganggap '2nd Year', '2nd', dan '2nd ' sebagai tiga kategori yang berbeda dan menghasilkan 9 kolom fitur terpecah yang merusak bobot model. Setelah diseragamkan menjadi 4 kategori baku, diperoleh distribusi merata: '2nd Year' (132), 'Final Year' (132), '3rd Year' (124), dan '1st Year' (112).\n\n"
        "3. Insight 3 (Pengaruh Pengalaman Proyek 'Projects' & Rating terhadap Kesenjangan Kompetensi):\n"
        "   • Temuan Data: Sebanyak 50.8% mahasiswa (254 dari 500) belum pernah mengerjakan proyek ('Projects' == 'No'), dan rata-rata 'Technical Rating' berada di angka 2.95 ± 1.40 (skala 1–5) serta 'SoftSkill Rating' 2.96 ± 1.44. Grafik 3 (Boxplot) menunjukkan bahwa mahasiswa tanpa pengalaman proyek ('No') yang memiliki Technical Rating <= 2 seluruhnya jatuh ke dalam kategori 'High_Gap' dengan median skor kesiapan di bawah 45/100.\n"
        "   • Dampak bagi Pemodelan: Interaksi antara 'Technical_Rating', 'SoftSkill_Rating', dan status 'Projects' memiliki korelasi negatif terkuat terhadap risiko 'High_Gap' (r = -0.68 pada Heatmap Grafik 4), sehingga sangat efektif dijadikan dasar pembuatan fitur turunan 'skor_kesiapan_industri'.\n\n"
        "4. Insight 4 (Distribusi Label & Pemeriksaan Class Imbalance serta Privasi PII):\n"
        "   • Temuan Data: Distribusi minat karier ('Career Interest') terbagi ke dalam 6 profesi dipimpin oleh 'Cloud Architect' (19.6% / 98 orang) dan 'Software Developer' (18.4% / 92 orang), sedangkan distribusi label kesenjangan kompetensi menunjukkan proporsi 'Ready' sebesar 61.8% (309 mahasiswa) berbanding 'High_Gap' sebesar 38.2% (191 mahasiswa). Selain itu, kolom 'Name' dan 'email_id' memiliki 500 nilai unik (100% berbeda tiap baris).\n"
        "   • Dampak bagi Pemodelan: Adanya ketimpangan kelas (61.8% vs 38.2%) mengharuskan penggunaan parameter stratify=y saat train_test_split dan penggunaan metrik F1-Score serta Recall. Sementara itu, kolom 'Name' dan 'email_id' wajib dibuang untuk melindungi privasi data (PII) dan mencegah overfitting."
    )
    add_box_section(doc, "Insight EDA ('DMA DATASET.xlsx'):", insight_eda_text, bg_hex="EFF6FF", border_title_rgb=RGBColor(0x1E, 0x3A, 0x8A))

    # SOAL 3
    add_sub_heading("Soal 3 – Pembersihan Data dan Feature Engineering ('DMA DATASET.xlsx')")
    add_box_section(
        doc,
        "Output Soal 3 – Data Sebelum dan Sesudah Dibersihkan serta Fitur Baru pada 'DMA DATASET.xlsx':",
        res["out_a5"] + "\n\nFitur Turunan Baru pada DMA DATASET:\n" + res["out_a6"] + "\n\nMatriks TF-IDF Teks Profil & Keahlian:\n" + res["out_a12_1"] + "\n\nFitur Time Series Harian (shift(1)):\n" + res["out_a12_2"],
        bg_hex="F8FAFC",
        is_code=True
    )
    alasan_pembersihan_text = (
        "1. Pembersihan Spasi Nama Kolom & Penggantian Nama Kolom Ambigu: Menghapus spasi di awal/akhir pada 'Year ', '     Programming Languages', ' Projects', 'Challenges ', dan ' Method', serta mengubah 'Rating' menjadi 'Technical_Rating' dan 'Rating.1' menjadi 'SoftSkill_Rating' agar skema data jelas.\n"
        "2. Penyeragaman 9 Variasi Penulisan Kolom 'Year' Menjadi 4 Tingkat Baku: Mengubah '1st ', '1st' -> '1st Year'; '2nd ', '2nd' -> '2nd Year'; dan 'Final ', 'Final' -> 'Final Year' agar OneHotEncoder tidak membentuk kolom kategori palsu.\n"
        "3. Penghapusan Kolom 'Unnamed: 13' (100% NaN) serta Kolom PII ('Name', 'email_id'): 'Unnamed: 13' dibuang karena seluruh 500 barisnya kosong, sedangkan 'Name' dan 'email_id' dibuang untuk menjaga privasi data mahasiswa serta mencegah model menghafal identitas individu.\n"
        "4. Pembuatan 3 Fitur Turunan Berbasis Pengetahuan Domain pada DMA DATASET:\n"
        "   • total_item_skill = jml_tech_skills + jml_prog_langs: Menghitung banyaknya variasi keahlian teknis dan bahasa pemrograman yang dikuasai mahasiswa dari teks koma-terpisah.\n"
        "   • skor_kesiapan_industri = 0.55*(Technical_Rating*20) + 0.30*(SoftSkill_Rating*20) + 0.15*(punya_proyek_biner*100): Mengukur skor kesiapan karier (skala 0–100) dengan menggabungkan rating teknis (55%), softskill (30%), dan kepemilikan proyek nyata (15%).\n"
        "   • indeks_defisit_skill = max(0, 80 - skor_kesiapan_industri) & Binning 'tingkat_kesiapan': Mengukur jarak kekurangan skor mahasiswa terhadap ambang kesiapan industri (80 poin) dan mengelompokkannya ke 'Butuh_Intensif', 'Sedang_Berkembang', dan 'Siap_Industri'.\n"
        "5. Ekstraksi Teks TF-IDF: Menggabungkan teks 'Current Course', 'Technical Skills', 'Programming Languages', 'Soft Skills', 'Career Interest', dan 'Challenges' dalam huruf kecil ke dalam TfidfVectorizer untuk menangkap pola kecocokan keahlian terhadap minat karier."
    )
    add_box_section(doc, "Alasan setiap keputusan pembersihan dan fitur baru:", alasan_pembersihan_text, bg_hex="EFF6FF", border_title_rgb=RGBColor(0x1E, 0x3A, 0x8A))

    # SOAL 4
    add_sub_heading("Soal 4 – Baseline Project ('DMA DATASET.xlsx')")
    add_box_section(
        doc,
        "Output Soal 4 – Tabel Perbandingan Baseline, TF-IDF + K-Means pada DMA DATASET, & RAG Retrieval:",
        res["out_a7"] + "\n\nTabel Perbandingan Kinerja Model pada Data Uji (DMA DATASET):\n" + res["out_a8"] + "\n\n5-Fold Cross-Validation:\n" + res["out_a9"] + "\n\n" + res["out_a11"] + "\n\nRetrieval TF-IDF (Recall@k & MRR):\n" + res["out_a12_3"],
        bg_hex="F8FAFC",
        is_code=True,
        img_path="docs/images/soal4_kmeans_baseline.png"
    )
    penjelasan_baseline_text = (
        "YA, Baseline Machine Learning di dalam Pipeline terbukti JAUH LEBIH BAIK dibandingkan Baseline Non-AI maupun Dummy Classifier pada 'DMA DATASET.xlsx', dengan bukti kuantitatif berikut:\n\n"
        "1. Keunggulan Pipeline Classification ML vs Baseline Non-AI (Daftar Kata Kunci Skill Manual):\n"
        "   • Dummy Classifier (menebak kelas terbanyak 'Ready') hanya memperoleh Accuracy = 0.620 namun Recall = 0.000 dan F1 = 0.000 karena sama sekali gagal mendeteksi mahasiswa berkategori High_Gap.\n"
        "   • Baseline Non-AI (Pencocokan Daftar Kata Kunci Skill Manual terhadap Career Interest) hanya mencapai Accuracy = 0.600, Recall = 0.289, dan F1-Score = 0.355. Hal ini terjadi karena banyak mahasiswa di DMA DATASET yang mencantumkan kata kunci bahasa pemrograman populer (seperti Python atau Java), tetapi memiliki Technical Rating rendah (1 atau 2) dan belum memiliki pengalaman proyek ('Projects' == 'No').\n"
        "   • Sebaliknya, Pipeline Machine Learning (ColumnTransformer + Model) melesat jauh melampaui baseline non-AI: Logistic Regression mencapai Accuracy = 0.950 dan F1 = 0.937 (5-Fold CV F1 = 0.884 ± 0.048), Random Forest mencapai Accuracy = 0.930 dan F1 = 0.899 dengan stabilitas lintas-fold terbaik (5-Fold CV F1 = 0.916 ± 0.018), dan Decision Tree mencapai Accuracy = 0.990 dan F1 = 0.987 (5-Fold CV F1 = 0.949 ± 0.030).\n\n"
        "2. Keberhasilan Pengolahan TF-IDF + K-Means pada DMA DATASET (Tabel Panduan No. 15):\n"
        "   • Evaluasi Silhouette Score dari k=2 (0.111) meningkat secara konsisten hingga mencapai nilai puncak tertinggi tepat pada k=6 (Silhouette Score = 0.294) sebelum turun kembali pada k=7 (0.253). Nilai k=6 secara alami bersesuaian dengan 6 kelompok minat karier ('Career Interest') di dalam DMA DATASET (Cluster 0: Cybersecurity Expert, Cluster 1: Cloud Architect, Cluster 2: Data Scientist, Cluster 3: Software Developer, Cluster 4: Business Analyst, Cluster 5: AI Engineer).\n\n"
        "3. Keberhasilan Baseline Retrieval Dokumen RAG Berbasis TF-IDF:\n"
        "   • Pengujian retrieval TF-IDF menghasilkan Recall@1 = 1.00, Recall@3 = 1.00, dan MRR = 1.00, serta mengembalikan daftar kosong [] untuk pertanyaan di luar domain."
    )
    add_box_section(doc, "Apakah baseline ML sudah lebih baik daripada baseline non-AI? Jelaskan:", penjelasan_baseline_text, bg_hex="EFF6FF", border_title_rgb=RGBColor(0x1E, 0x3A, 0x8A))

    # SOAL 5
    add_sub_heading("Soal 5 – Pemeriksaan Data Leakage ('DMA DATASET.xlsx')")
    hasil_leakage_text = (
        "a. Daftar Fitur yang Dipakai & Identifikasi Fitur Penyebab Target Leakage:\n"
        "   • Fitur Valid yang Dipakai pada Pipeline X (11 Fitur Input Asesmen): 'Technical_Rating', 'SoftSkill_Rating', 'jml_tech_skills', 'jml_prog_langs', 'total_item_skill', 'Year', 'Current Course', 'Projects', 'Career Interest', 'Method', dan 'teks_profil_lengkap' (TF-IDF).\n"
        "   • Fitur Penyebab Target Leakage yang Dikeluarkan dari X: Kolom turunan pasca-evaluasi seperti 'tingkat_kesiapan', 'skor_kesiapan_industri', dan 'indeks_defisit_skill' (serta 'Support required' yang merupakan rekomendasi intervensi akhir) TIDAK dimasukkan ke dalam matriks fitur X pada pelatihan klasifikasi karena diturunkan dari formula penentu status gap. Selain itu, kolom identitas 'Name' dan 'email_id' telah dibuang.\n\n"
        "b. Pencegahan Train-Test Contamination melalui Pipeline & ColumnTransformer:\n"
        "   • Dataset (500 baris) di-split terlebih dahulu menggunakan train_test_split(X, y, test_size=0.2, stratify=y, random_state=42) menjadi 400 baris data latih dan 100 baris data uji.\n"
        "   • Seluruh proses SimpleImputer (median & most_frequent), StandardScaler, OneHotEncoder, dan TfidfVectorizer dibungkus di dalam ColumnTransformer dan Pipeline scikit-learn, sehingga perhitungan rata-rata, simpangan baku, dan bobot IDF murni dihitung dari 400 baris data latih (X_train).\n\n"
        "c. Pemeriksaan Overlap Objek / Mahasiswa yang Sama di Data Latih dan Data Uji:\n"
        "   • Pengujian irisan identitas unik len(set(id_train).intersection(set(id_test))) pada 'email_id' membuktikan terdapat tepat 0 mahasiswa yang muncul di data latih dan data uji sekaligus."
    )
    add_box_section(doc, "Hasil pemeriksaan data leakage dan perbaikannya:", hasil_leakage_text, bg_hex="EFF6FF", border_title_rgb=RGBColor(0x1E, 0x3A, 0x8A))

    # ==========================================================================
    # LEMBAR PENGESAHAN
    # ==========================================================================
    doc.add_paragraph().paragraph_format.space_before = Pt(8)
    tbl_sign = doc.add_table(rows=4, cols=2)
    tbl_sign.alignment = WD_TABLE_ALIGNMENT.CENTER

    r0_c0, r0_c1 = tbl_sign.rows[0].cells[0], tbl_sign.rows[0].cells[1]
    r0_c0.width, r0_c1.width = Inches(3.25), Inches(3.25)
    set_cell_background(r0_c0, "F1F5F9")
    set_cell_background(r0_c1, "F1F5F9")
    p_s0 = r0_c0.paragraphs[0]
    p_s0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_s0.add_run("Mahasiswa").bold = True
    p_s1 = r0_c1.paragraphs[0]
    p_s1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_s1.add_run("Asisten Dosen").bold = True

    r1_c0, r1_c1 = tbl_sign.rows[1].cells[0], tbl_sign.rows[1].cells[1]
    r1_c0.width, r1_c1.width = Inches(3.25), Inches(3.25)
    set_cell_margins(r1_c0, top=120, bottom=80, left=100, right=100)
    set_cell_margins(r1_c1, top=120, bottom=80, left=100, right=100)
    p_mhs = r1_c0.paragraphs[0]
    p_mhs.add_run("\n\n1. Julia Rahma Nurmawati (NIM. V3925009)\n2. Aisyah Nurul Ilmi Prianto (NIM. V3925018)").font.size = Pt(9.5)
    p_asdos = r1_c1.paragraphs[0]
    p_asdos.add_run("\n\nNama Asisten Dosen: ...........................................\nNIM: .....................................................................").font.size = Pt(9.5)

    r2_c0 = tbl_sign.rows[2].cells[0].merge(tbl_sign.rows[2].cells[1])
    set_cell_background(r2_c0, "F8FAFC")
    p_men = r2_c0.paragraphs[0]
    p_men.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_men = p_men.add_run("Mengetahui\nDosen Mata Kuliah")
    r_men.bold = True
    r_men.font.size = Pt(9.5)

    r3_c0 = tbl_sign.rows[3].cells[0].merge(tbl_sign.rows[3].cells[1])
    set_cell_margins(r3_c0, top=140, bottom=80, left=100, right=100)
    p_dos = r3_c0.paragraphs[0]
    p_dos.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_dos = p_dos.add_run("\n\nMasbahah, S.Pd., M.Pd.\nNIP. 1987052520200801")
    r_dos.bold = True
    r_dos.font.size = Pt(10)

    os.makedirs(os.path.dirname(output_docx), exist_ok=True)
    doc.save(output_docx)
    root_copy = "LAPORAN_TUGAS_MANDIRI_HARI1_SKILLGAP_AI.docx"
    shutil.copyfile(output_docx, root_copy)
    print(f"Laporan Word berbasis DMA DATASET.xlsx berhasil diperbarui di:\n1) {output_docx}\n2) {root_copy}")

if __name__ == "__main__":
    with open("data/processed/praktikum_outputs.json", "r", encoding="utf-8") as f:
        res_data = json.load(f)
    create_word_report(res_data)
